#!/usr/bin/env python3
"""build_review_pptx.py — the MOVABLE per-biome deck as a real .pptx.

Same verified data as build_review_deck.py (imported, not re-derived), rendered
so the owner can DRAG creatures between biome slides. Each creature is an
individual picture sized by true drawSize; its label rides the shape name and
its intention/decision/size ride the shape's alt-text (Selection pane / right-
click → Edit Alt Text / Description). Plants sit on a lower band; two text boxes
carry the biome text and BENCH's comments. Grouping slides follow the biomes.
"""
from __future__ import annotations
import io, sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_review_deck as B

EMU_IN = 914400
SLIDE_W, SLIDE_H = 13.333, 7.5
BG = RGBColor(0x22,0x1a,0x12); PANEL = RGBColor(0x2e,0x23,0x18)
INK = RGBColor(0xf0,0xe3,0xd0); DIM = RGBColor(0xb9,0xa4,0x88)
ACCENT = RGBColor(0xc9,0x8a,0x3e); INCOL = RGBColor(0x7f,0xae,0x6a); ARRCOL = RGBColor(0x5b,0x9b,0xd5)

K = 0.62          # inches of picture height per drawSize cell
HMIN, HMAX = 0.32, 2.4
ANIM_TOP, ANIM_BOT = 1.02, 4.05   # animal band (inches)
PLANT_TOP, PLANT_BOT = 4.18, 5.02
PLANT_H = 0.6
GAP = 0.14        # inches between critters
CAP_H = 0.16

TMP = Path("/tmp/claude-1000/deck_thumbs"); TMP.mkdir(parents=True, exist_ok=True)
_thumb: dict[str, Path | None] = {}
def thumb(defName, folder: Path):
    if defName in _thumb: return _thumb[defName]
    src = folder / f"{defName}.png"; out = None
    if src.exists():
        try:
            im = Image.open(src).convert("RGBA"); im.thumbnail((180,180), Image.LANCZOS)
            out = TMP / f"{folder.name}_{defName}.png"; im.save(out, "PNG", optimize=True)
        except Exception: out = None
    _thumb[defName] = out; return out

def dims(defName, folder):
    p = thumb(defName, folder)
    if not p: return None, 1.0
    with Image.open(p) as im: return p, im.width/im.height

def set_alt(shape, name, descr):
    el = shape._element.nvPicPr.cNvPr if shape.shape_type == 13 else shape._element._nvGrpSpPr.cNvPr \
         if False else None
    try:
        cNvPr = shape._element._nvXxPr.cNvPr
    except AttributeError:
        cNvPr = shape._element.getchildren()[0].getchildren()[0]
    cNvPr.set("name", name[:120]); cNvPr.set("descr", descr[:900])

def hpx(ds): return max(HMIN, min(HMAX, K*(ds or 1)))

def bg(slide):
    r = slide.background.fill; r.solid(); r.fore_color.rgb = BG

def textbox(slide, x, y, w, h, title, body, tcol):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; run = p.add_run(); run.text = title
    run.font.size = Pt(9); run.font.bold = True; run.font.color.rgb = tcol
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = body[:1500]
    r2.font.size = Pt(8); r2.font.color.rgb = DIM
    tb.fill.solid(); tb.fill.fore_color.rgb = PANEL
    tb.line.color.rgb = ACCENT; tb.line.width = Pt(0.5)
    return tb

def lay_row(slide, items, folder, top, bottom, fixed_h=None):
    """Greedy row-pack items into [top,bottom]; uniform down-scale to fit."""
    band = bottom - top
    def size_of(it):
        if fixed_h:
            _, ar = dims(it["defName"], folder); return fixed_h, fixed_h*max(ar,0.4)
        h = hpx(it.get("drawSize",1)); _, ar = dims(it["defName"], folder)
        return h, h*max(ar,0.4)
    # first pass rows at scale 1
    def pack(scale):
        rows=[[]]; x=0.3; used=[0.0]; rowH=[0.0]
        for it in items:
            h,w = size_of(it); h*=scale; w*=scale
            if x+w > SLIDE_W-0.3 and rows[-1]:
                rows.append([]); x=0.3; used.append(0.0); rowH.append(0.0)
            rows[-1].append((it,x,h,w)); x+=w+GAP; rowH[-1]=max(rowH[-1],h)
        total = sum(rh+CAP_H+0.08 for rh in rowH)
        return rows, rowH, total
    scale=1.0
    rows,rowH,total = pack(scale)
    if total > band and total>0:
        scale = max(0.4, band/total); rows,rowH,total = pack(scale)
    y = top
    for ri,row in enumerate(rows):
        base = y + rowH[ri]
        for it,x,h,w in row:
            p = thumb(it["defName"], folder)
            note = it.get("note","")
            dec = it.get("decision"); origin=it.get("origin")
            meta = []
            if it.get("drawSize"): meta.append(f"drawSize {it['drawSize']}")
            if dec=="arrived": meta.append(f"moved in from {origin}")
            elif dec=="in": meta.append("assigned here")
            if it.get("art"): meta.append(f"art:{it['art']}")
            if it.get("conf"): meta.append(str(it["conf"]))
            descr = it.get("label","")+" ("+it["defName"]+")\n"+" · ".join(meta)+("\n“"+note+"”" if note else "")
            yy = base-h
            if p:
                pic = slide.shapes.add_picture(str(p), Inches(x), Inches(yy), height=Inches(h))
                set_alt(pic, it.get("label",it["defName"]), descr)
                dot = it.get("decision")
                if dot in ("in","arrived"):
                    d = slide.shapes.add_shape(9, Inches(x+w-0.09), Inches(yy-0.02), Inches(0.08),Inches(0.08))
                    d.fill.solid(); d.fill.fore_color.rgb = INCOL if dot=="in" else ARRCOL
                    d.line.fill.background()
            else:
                box = slide.shapes.add_shape(5, Inches(x), Inches(yy), Inches(max(w,0.7)), Inches(h))
                box.fill.solid(); box.fill.fore_color.rgb = PANEL; box.line.color.rgb=DIM; box.line.width=Pt(0.5)
                tf=box.text_frame; tf.word_wrap=True; r=tf.paragraphs[0].add_run(); r.text=it.get("label",it["defName"])
                r.font.size=Pt(6); r.font.color.rgb=DIM
                set_alt(box, it.get("label",it["defName"]), descr+"\n[no cached sprite]")
            # caption
            cap = slide.shapes.add_textbox(Inches(x), Inches(base+0.01), Inches(max(w,0.7)), Inches(CAP_H))
            cr = cap.text_frame.paragraphs[0].add_run(); cr.text = it.get("label",it["defName"])[:16]
            cr.font.size=Pt(6); cr.font.color.rgb=DIM; cap.text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER
        y = base + CAP_H + 0.1

def title(slide, txt, sub):
    tb = slide.shapes.add_textbox(Inches(0.3), Inches(0.18), Inches(12.7), Inches(0.7))
    tf=tb.text_frame; p=tf.paragraphs[0]; r=p.add_run(); r.text=txt
    r.font.size=Pt(22); r.font.bold=True; r.font.color.rgb=ACCENT
    r2=p.add_run(); r2.text="   "+sub; r2.font.size=Pt(11); r2.font.color.rgb=DIM

def label_band(slide, y, txt):
    tb=slide.shapes.add_textbox(Inches(0.3), Inches(y), Inches(6), Inches(0.2))
    r=tb.text_frame.paragraphs[0].add_run(); r.text=txt; r.font.size=Pt(8); r.font.color.rgb=DIM

def main():
    fauna = B.json.load(open(B.REVIEW/"round2/decisions_propagated.json"))["decisions"]
    flora = B.json.load(open(B.REVIEW/"flora_assignment_register.decisions.json"))["decisions"]
    census = B.load_census(); moves = B.load_moves()
    animals, plants = B.build_biomes(fauna, flora, moves, census)
    findings = B.load_findings(); groups = B.load_groups(census)
    fac = B.build_faction_group(census)
    if fac: groups["9. faction-territory fauna (Hutt / Helix / Wildsteam / Moisture Farmers)"]=fac

    prs = Presentation(); prs.slide_width=Inches(SLIDE_W); prs.slide_height=Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    order = sorted(animals, key=lambda b: -(len(animals.get(b,[]))+len(plants.get(b,[]))))
    for b in order:
        s = prs.slides.add_slide(blank); bg(s)
        aa = sorted(animals.get(b,[]), key=lambda a:-a.get("drawSize",1))
        pp = plants.get(b,[])
        title(s, b, f"{len(aa)} animals · {len(pp)} plants")
        label_band(s, ANIM_TOP-0.22, "ANIMALS — sized by drawSize · drag to move · hover/alt-text = intent")
        lay_row(s, aa, B.FAUNA_SPRITES, ANIM_TOP, ANIM_BOT)
        label_band(s, PLANT_TOP-0.2, "PLANTS")
        lay_row(s, pp, B.FLORA_SPRITES, PLANT_TOP, PLANT_BOT, fixed_h=PLANT_H)
        textbox(s, 0.3, 5.2, 6.3, 2.1, "BIOME TEXT (for review)", B.biome_text(b), DIM)
        textbox(s, 6.75, 5.2, 6.28, 2.1, "BENCH PER-BIOME COMMENTS", findings.get(b,"—"), ACCENT)

    # grouping slides
    for g, rows in groups.items():
        s = prs.slides.add_slide(blank); bg(s)
        title(s, g.split("(")[0].strip(), f"{len(rows)} creatures · grouping sheet")
        label_band(s, ANIM_TOP-0.22, g)
        lay_row(s, sorted(rows, key=lambda r:-r.get("drawSize",1)), B.FAUNA_SPRITES, ANIM_TOP, 6.9)

    out = B.R2 / "fauna_review_deck.pptx"
    prs.save(str(out))
    print(f"slides={len(prs.slides._sldIdLst)} size={out.stat().st_size/1e6:.2f}MB -> {out}")

if __name__ == "__main__":
    main()
